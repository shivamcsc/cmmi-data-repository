"""
Workbook parsing and loading.

parse_workbook() reads every worksheet, works out where the header row is and
which family the sheet belongs to, and returns a plain description of what it
found. load_upload() then writes one database row per sheet row.

Sheets that are not data collection sheets -- the estimation and size tabs
bundled into the CCTNS and CIBIL workbooks, revision-history tabs, scratch
tabs -- are recorded with a skip_reason rather than parsed.
"""
import datetime as dt
import hashlib
import re

from openpyxl.utils import get_column_letter

from sqlalchemy import Date, Float, String, select

from app import metrics_catalog as mc
from app.models import RECORD_MODELS, SheetTab, UploadBatch

HEADER_SEARCH_ROWS = 15
MIN_HEADER_MATCHES = 3
BLANK_RUN_LIMIT = 40
MAX_ROWS = 50000

#: Day first throughout -- these are Indian sheets, so 06-01-2026 is 6 January.
#: The spread is not carelessness on the teams' part: a date typed into a cell
#: Excel has formatted as text stays text, and each sheet was typed by hand, so
#: the same column holds several spellings. Every format below was found in the
#: files under storage/uploads; see scripts/scan_dates.py, which lists any that
#: still fail to parse.
_DATE_FORMATS = [
    "%d/%m/%Y, %I:%M %p", "%d/%m/%Y %I:%M %p", "%d/%m/%Y %H:%M", "%d/%m/%Y",
    "%d-%m-%Y %H:%M", "%d-%m-%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d",
    "%d.%m.%Y", "%d %b %Y", "%d-%b-%Y", "%b %d, %Y", "%m/%d/%Y",

    # CIBIL writes the whole stamp as text with a colon before the time, so the
    # cell reads 06-01-2026:10:00:00 however the column is formatted.
    "%d-%m-%Y:%H:%M:%S", "%d-%m-%Y:%H:%M", "%d/%m/%Y:%H:%M:%S", "%d/%m/%Y:%H:%M",
    "%d-%m-%Y %H:%M:%S", "%d/%m/%Y %H:%M:%S",

    # Two-digit years: 31-Mar-26. %y reads 26 as 2026.
    "%d-%b-%y", "%d %b %y", "%d/%b/%y", "%d-%b-%Y %H:%M", "%d %b %Y %H:%M",

    # A comma between date and time, with the space around am/pm going missing
    # either side of it: '15-06-2026, 2:33 pm', '18/08/2025, 2:20pm',
    # '01/09/2025,10:30 am'.
    "%d-%m-%Y, %I:%M %p", "%d-%m-%Y, %I:%M%p", "%d-%m-%Y,%I:%M %p",
    "%d/%m/%Y, %I:%M%p", "%d/%m/%Y,%I:%M %p", "%d/%m/%Y,%I:%M%p",
    "%d-%m-%Y, %H:%M", "%d/%m/%Y, %H:%M", "%d-%m-%Y,%H:%M", "%d/%m/%Y,%H:%M",
]
_NUMERIC = re.compile(r"^-?[\d,]*\.?\d+%?$")


class IngestError(Exception):
    pass


class DuplicateUpload(IngestError):
    """The identical file was already loaded for this project and period."""


# --------------------------------------------------------------------------
# Value coercion
# --------------------------------------------------------------------------
def to_date(value):
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    for fmt in _DATE_FORMATS:
        try:
            return dt.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def to_number(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if value is None:
        return None
    text = str(value).strip()
    if not text or not _NUMERIC.match(text):
        return None
    try:
        return float(text.rstrip("%").replace(",", ""))
    except ValueError:
        return None


def is_blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def as_text(value):
    if isinstance(value, (dt.date, dt.datetime)):
        return value.isoformat()
    return str(value).strip()


def column_types(model):
    """{column: (kind, max_length)} so values are coerced and truncated to fit."""
    out = {}
    for column in model.__table__.columns:
        kind = column.type
        if isinstance(kind, Date):
            out[column.key] = ("date", None)
        elif isinstance(kind, Float):
            out[column.key] = ("number", None)
        elif isinstance(kind, String):
            out[column.key] = ("text", kind.length)
        else:
            out[column.key] = ("text", None)
    return out


def coerce(value, kind, limit):
    """Fit a cell to its column. A value that will not convert is dropped, not guessed."""
    if is_blank(value):
        return None
    if kind == "date":
        return to_date(value)
    if kind == "number":
        return to_number(value)
    text = as_text(value)
    return text[:limit] if limit else text


# --------------------------------------------------------------------------
# Reading workbooks
# --------------------------------------------------------------------------
#: columns whose value is written once and understood to hold until it changes.
#: A blank effort or date means blank; a blank project name does not.
CARRY_DOWN = {
    "development": ("department", "project_name", "service_name"),
    "services": (),
}


def merged_spans(worksheet):
    """
    {(row, column): value} for every cell a merge covers but the file leaves empty.

    Excel stores a merged block's value in its top-left cell only, so a module
    merged down its three sub-module rows reads as the name once and then None
    twice. The reader fills the rest back in.
    """
    spans = {}
    for area in worksheet.merged_cells.ranges:
        # Down only, never across. A merge spanning one row is a banner or a
        # note -- 'Note: External Testing is ongoing' is merged B9:J9, nine
        # columns wide -- and spreading it sideways would write that sentence
        # into the size and effort columns.
        if area.max_row == area.min_row:
            continue
        value = worksheet.cell(area.min_row, area.min_col).value
        if value is None:
            continue
        for row in range(area.min_row + 1, area.max_row + 1):
            spans[(row, area.min_col)] = value
    return spans


def fill_merged(row, number, spans):
    """Put the merged value back into the cells Excel left empty."""
    if not spans:
        return row
    for position in range(len(row)):
        if row[position] is None:
            value = spans.get((number, position + 1))
            if value is not None:
                row[position] = value
    return row


def has_measurement(values, positions):
    """
    True when the row carries at least one measured value.

    A data collection row records something: a size, an effort, a date, a
    defect count. The blocks the teams write below the table -- a note typed
    into the Department cell, a Legend of 'Weekends excluded' and '1 working
    day = 8 hours' -- carry text and nothing else, and the reader would
    otherwise take them as modules, because they sit fewer than
    BLANK_RUN_LIMIT rows under the last real one.
    """
    for position in positions:
        if position < len(values) and not is_blank(values[position]):
            return True
    return False


def measure_positions(columns, family):
    """Where the columns that hold a measurement sit, identity columns aside."""
    identity = set(CARRY_DOWN.get(family, ()))
    return [i for i, (column, _raw, _ignored) in enumerate(columns)
            if column and column not in identity]


def carry_down(rows, columns, family):
    """
    Carry an identity column's value down the rows beneath it.

    'NFDP' is typed once at the top of Project Name and the cells below are
    simply left empty -- no merge, so there is nothing to expand. Every row
    belongs to that project all the same. Only the columns in CARRY_DOWN are
    treated this way: a blank effort or date means blank.
    """
    wanted = CARRY_DOWN.get(family, ())
    if not wanted:
        return
    positions = [(i, c) for i, (c, _raw, _ignored) in enumerate(columns) if c in wanted]
    if not positions:
        return
    last = {}
    for row in rows:
        values = row["values"]
        for position, column in positions:
            if position >= len(values):
                continue
            if is_blank(values[position]):
                if column in last:
                    values[position] = last[column]
            else:
                last[column] = values[position]


def read_sheets(path):
    """[(sheet_name, [[cell, ...], ...]), ...] for xlsx/xlsm/xls/csv."""
    name = str(path).lower()
    if name.endswith((".xlsx", ".xlsm")):
        import openpyxl

        # read_only would be lighter, but merged_cells is not populated in that
        # mode and the teams merge the module column down its sub-module rows,
        # so the module name would be read for the first row only.
        book = openpyxl.load_workbook(path, data_only=True)
        try:
            out = []
            for sheet_name in book.sheetnames:
                worksheet = book[sheet_name]
                spans = merged_spans(worksheet)
                rows, blank_run = [], 0
                for number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                    row = fill_merged(list(row), number, spans)
                    if all(is_blank(c) for c in row):
                        blank_run += 1
                        rows.append(row)
                        if blank_run >= BLANK_RUN_LIMIT:
                            break
                    else:
                        blank_run = 0
                        rows.append(row)
                    if len(rows) >= MAX_ROWS:
                        break
                out.append((sheet_name, rows))
            return out
        finally:
            book.close()

    import pandas as pd

    if name.endswith(".csv"):
        frame = pd.read_csv(path, header=None, dtype=object,
                            keep_default_na=False, na_filter=False)
        return [("CSV", frame.values.tolist())]
    book = pd.read_excel(path, sheet_name=None, header=None, dtype=object)
    return [(n, f.where(f.notna(), None).values.tolist()) for n, f in book.items()]


def find_header_row(rows):
    """
    Locate the header row and its family.

    Header position varies by project -- row 1 in the CCTNS phase tabs, row 3 in
    CIBIL, row 5 in Digipay and Sales Force CRM, where banner rows sit above it.
    The row resolving the most catalogue columns wins.
    """
    best = (None, None, 0)
    for index, row in enumerate(rows[:HEADER_SEARCH_ROWS]):
        if len([c for c in row if not is_blank(c)]) < MIN_HEADER_MATCHES:
            continue
        family, scores = mc.detect_family(row)
        if family is not None and scores[family] > best[2]:
            best = (index, family, scores[family])
    return best


def parse_workbook(path):
    """Describe every worksheet: its columns, family, and the rows to load."""
    parsed = []
    for order, (sheet_name, rows) in enumerate(read_sheets(path)):
        info = {
            "sheet_name": sheet_name, "sheet_index": order, "header_row": None,
            "data_start_row": None, "family": None, "columns": [], "rows": [],
            "row_count": 0, "column_count": 0, "parsed": False, "skip_reason": None,
        }
        if not rows:
            info["skip_reason"] = "worksheet is empty"
            parsed.append(info)
            continue

        header_index, family, _score = find_header_row(rows)
        if header_index is None:
            info["skip_reason"] = "no data collection header row found"
            info["column_count"] = max((len(r) for r in rows), default=0)
            parsed.append(info)
            continue

        header_row = rows[header_index]
        info.update({
            "header_row": header_index + 1,
            "data_start_row": header_index + 2,
            "family": family,
            "columns": mc.resolve_headers(header_row, family),
            "column_count": len(header_row),
            "parsed": True,
        })

        for offset, row in enumerate(rows[header_index + 1:], start=header_index + 2):
            if all(is_blank(c) for c in row):
                continue
            info["rows"].append({"row_index": offset, "values": list(row)})

        # Drop the notes and legend blocks before carrying anything down: a
        # note typed into the Department cell below the table would otherwise
        # be carried onto every row beneath it.
        measures = measure_positions(info["columns"], family)
        if measures:
            info["rows"] = [r for r in info["rows"] if has_measurement(r["values"], measures)]
        carry_down(info["rows"], info["columns"], family)
        info["row_count"] = len(info["rows"])
        parsed.append(info)
    return parsed


# --------------------------------------------------------------------------
# Loading
# --------------------------------------------------------------------------
def file_digest(path):
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_upload(session, upload, stored_path, allow_duplicate=False):
    """Parse the stored file and write its rows against an UploadBatch."""
    digest = file_digest(stored_path)
    # Any earlier load of the same bytes counts, including one that failed to
    # parse -- otherwise an unreadable file is silently re-added every import.
    if not allow_duplicate:
        clash = session.scalar(
            select(UploadBatch).where(
                UploadBatch.file_hash == digest,
                UploadBatch.project_id == upload.project_id,
                UploadBatch.period_id == upload.period_id,
                UploadBatch.id != upload.id,
            )
        )
        if clash is not None:
            raise DuplicateUpload(
                "This exact file is already loaded for this project and period "
                "(upload #%d, %s)." % (clash.id, clash.uploaded_at.strftime("%d %b %Y %H:%M"))
            )
    upload.file_hash = digest

    sheets = parse_workbook(stored_path)
    total_rows = 0

    for info in sheets:
        tab = SheetTab(
            upload_id=upload.id,
            sheet_name=str(info["sheet_name"])[:255],
            sheet_index=info["sheet_index"],
            sheet_kind=info["family"] or "other",
            header_row=info["header_row"],
            data_start_row=info["data_start_row"],
            row_count=info["row_count"],
            column_count=info["column_count"],
            was_parsed=info["parsed"],
            skip_reason=info["skip_reason"],
            column_map=({column: get_column_letter(position + 1)
                         for position, (column, _raw, _ig) in enumerate(info["columns"])
                         if column} or None),
        )
        session.add(tab)
        session.flush()
        if not info["parsed"]:
            continue

        family = info["family"]
        model = RECORD_MODELS[family]
        types = column_types(model)

        for record in info["rows"]:
            values, extras = {}, {}
            for position, (column, raw_header, ignored) in enumerate(info["columns"]):
                if ignored or is_blank(raw_header):
                    continue
                cell = record["values"][position] if position < len(record["values"]) else None
                if is_blank(cell):
                    continue
                if column is None:
                    extras[str(raw_header)[:200]] = as_text(cell)[:500]
                    continue
                kind, limit = types.get(column, ("text", None))
                coerced = coerce(cell, kind, limit)
                if coerced is not None:
                    values[column] = coerced

            if not values and not extras:
                continue

            session.add(model(
                upload_id=upload.id,
                sheet_id=tab.id,
                category_id=upload.category_id,
                project_id=upload.project_id,
                period_id=upload.period_id,
                upload_name=upload.original_filename[:512],
                period_label=upload.period.label[:128],
                sheet_name=str(info["sheet_name"])[:255],
                row_index=record["row_index"],
                extra_columns=extras or None,
                **values
            ))
            total_rows += 1

    upload.sheet_count = sum(1 for s in sheets if s["parsed"])
    upload.row_count = total_rows
    if total_rows == 0:
        upload.status = "failed"
        upload.error_message = "No data collection rows were found in this workbook."
    elif any(not s["parsed"] for s in sheets):
        upload.status = "partial"
    else:
        upload.status = "success"

    # Uploads stand on their own. A later sheet for the same project and period
    # no longer marks the earlier one superseded: the period label is inferred
    # from folder text written by eleven different teams, so "same period" is
    # not a safe enough test for "replaces". Only a rejection sets the flag now.

    session.commit()
    return sheets
